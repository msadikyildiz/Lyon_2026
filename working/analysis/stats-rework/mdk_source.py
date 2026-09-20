"""Export MDK records with confirmed zero counts and explicit detection thresholds."""
import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import median_abs_deviation
from manifest import HERE, ROOT, MDK


def read_records(figure):
    panel, rel, stem, note = MDK[figure]
    path = ROOT / 'working/figures' / rel
    d = pd.read_excel(path, sheet_name='Sheet1')
    d = d.rename(columns={'COUNT': 'Count', 'DILUTION': 'Dilution'})
    d['source_excel_row'] = d.index + 2
    d = d[d.Name.notna()].copy()
    d['count_effective'] = d.Count
    d['dilution_effective'] = d.Dilution
    d['factor_effective'] = d.Factor
    d['clarification'] = ''
    d['possible_handling_failure'] = False
    if figure == 'Figure 3':
        for row, day, hour in [(62, 1, 7), (146, 2, 3)]:
            mask = d.source_excel_row.eq(row)
            record = d.loc[mask].iloc[0]
            assert (record.Day, record.Name, record.Replicate, record.Time) == (day, 'P1', 1, hour)
            assert pd.isna(record.Count) and pd.isna(record.Dilution)
            d.loc[mask, ['count_effective', 'dilution_effective', 'factor_effective']] = [0, -1, 10]
            d.loc[mask, 'clarification'] = 'Confirmed 19 September 2026: zero colonies, tenfold-concentrated sample; original blank cells retained'
            d.loc[mask, 'possible_handling_failure'] = row == 146
    d['count_status'] = np.where(d.count_effective.isna(), 'missing',
                                np.where(d.count_effective == 0, 'observed zero; below detection', 'measured'))
    d['fraction_workbook'] = d.Fraction
    d['fraction_observed'] = d.Fraction.where(d.Count.notna() & (d.Count > 0))
    d['included_in_original_panel'] = True
    if figure == 'Figure 2':
        d.loc[d.Day == 1, 'included_in_original_panel'] = False
    d['replicate_original'] = d.Replicate
    if figure in ['Figure 2', 'Figure 3']:
        day = 3 if figure == 'Figure 2' else 2
        d.loc[(d.Day == day) & (d.Name == 'P1'), 'Replicate'] += 3
    ws = load_workbook(path, data_only=False, read_only=True)['Sheet1']
    d['cfu_formula'] = [ws[f'K{r}'].value for r in d.source_excel_row]
    d['formula_volume_mL'] = np.where((d.Count > 0) & (d.CFU > 0), d.Count / d.Factor / d.CFU, np.nan)
    zeros = d.count_effective.eq(0)
    for i, r in d.loc[zeros].iterrows():
        match = re.fullmatch(r'=H\d+/\(J\d+\*(0\.\d+)\)', r.cfu_formula)
        if match is None:
            raise ValueError(f'Unrecognized zero-count volume formula: {r.cfu_formula}')
        d.loc[i, 'formula_volume_mL'] = float(match[1])
    d['below_detection'] = zeros
    d['detection_limit'] = 1 / (d.factor_effective * d.formula_volume_mL * d.TimeZero)
    d['fraction_lower'] = d.fraction_observed.where(~zeros, 0)
    d['fraction_upper'] = d.fraction_observed.where(~zeros, d.detection_limit)
    d['fraction_plot'] = d.fraction_upper
    d['record_note'] = note
    d['source_workbook'] = str(path.relative_to(ROOT))
    d['figure'] = figure
    return d


def main():
    for figure, (_, _, stem, _) in MDK.items():
        if figure == 'Supplementary Figure 11':
            continue
        d = read_records(figure)
        selected = d[d.included_in_original_panel]
        summaries = []
        for (name, time), g in selected.groupby(['Name', 'Time']):
            v = g.fraction_observed.dropna()
            lo, hi = g.fraction_lower.dropna(), g.fraction_upper.dropna()
            n_zero = int(g.below_detection.sum())
            n_missing = int(g.count_effective.isna().sum())
            assert len(v) + n_zero + n_missing == len(g)
            if n_zero:
                assert n_zero == 1 and g.loc[g.below_detection, 'fraction_upper'].max() < v.min()
            medians = [float(x.median()) for x in (lo, hi)]
            if not np.isclose(*medians, rtol=1e-12, atol=0):
                raise ValueError('Median needs an interval for this censoring configuration')
            spreads = [median_abs_deviation(x, scale='normal') for x in (lo, hi)]
            summaries.append(dict(Name=name, Time=time, n_records=len(g), n_observed=len(v),
                                  n_zero=n_zero, n_missing=n_missing, median=medians[0],
                                  median_lower=medians[0], median_upper=medians[1],
                                  mad_scaled=max(spreads), mad_scaled_min=min(spreads), mad_scaled_max=max(spreads),
                                  median_workbook=g.Fraction.median(),
                                  mad_scaled_workbook=median_abs_deviation(g.Fraction, scale='normal'),
                                  note='Zero-colony records retained as below detection; median checked over censoring endpoints; MAD normal-scaled'))
        s = pd.DataFrame(summaries)
        if figure in ['Figure 5', 'Supplementary Figure 10']:
            if (s.n_records != 3).any() or s.n_missing.any():
                raise ValueError(f'{figure}: expected three observed biological replicates')
            if not np.allclose(d.formula_volume_mL, .02, rtol=1e-12, atol=0):
                raise ValueError(f'{figure}: unexpected volume conversion')
        d.to_csv(HERE / 'out' / f'{stem}_records.csv', index=False)
        s.to_csv(HERE / 'out' / f'{stem}_summary.csv', index=False)
        if figure == 'Figure 3':
            subset = selected[(selected.Name == 'P1') & (selected.Time == 3)]
            without = subset.loc[~subset.possible_handling_failure, 'fraction_observed'].dropna()
            primary = s[(s.Name == 'P1') & (s.Time == 3)].iloc[0]
            pd.DataFrame([
                dict(scenario='primary; retain confirmed zero with possible handling failure', n=len(subset), median=primary['median'], mad_scaled=primary.mad_scaled),
                dict(scenario='omit day 2 replicate 1 at 3 h', n=len(without), median=without.median(), mad_scaled=median_abs_deviation(without, scale='normal'))
            ]).to_csv(HERE / 'out/Fig3e_handling_sensitivity.csv', index=False)
        print(f'{figure}: {len(d)} records; {selected.below_detection.sum()} zeros; {selected.count_effective.isna().sum()} missing')


if __name__ == '__main__':
    main()
