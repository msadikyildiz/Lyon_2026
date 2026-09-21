"""Assemble the Source Data workbook with explicit coverage and provenance."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from manifest import HERE, ROOT, panel_manifest, FIT_PANELS, MDK
from stats import split_strain
from manifest import comparison_cohort

DEST = ROOT / 'working/source-data'
SOURCE = ROOT / 'data/supplementary_tables.json'
COVERAGE = []
EXPECTED_CELLS = []


def sheet_name(figure):
    return figure.replace('Supplementary Figure ', 'Supp Fig ').replace('Supplementary Table ', 'Supp Table ')


def block(ws, title, table, source, note=''):
    row = ws.max_row + 2 if ws.max_row > 1 else 1
    ws.cell(row, 1, title).font = Font(name='Aptos', bold=True, size=12, color='16324F')
    ws.cell(row + 1, 1, 'Source: ' + str(source))
    if note:
        ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 2, end_column=8)
        ws.cell(row + 2, 1, note).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row + 2].height = 36
    header = row + 4
    for j, col in enumerate(table.columns, 1):
        c = ws.cell(header, j, str(col))
        c.fill = PatternFill('solid', fgColor='16324F')
        c.font = Font(name='Aptos', bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True)
    for i, values in enumerate(table.itertuples(index=False, name=None), header + 1):
        for j, value in enumerate(values, 1):
            if isinstance(value, (list, tuple, dict, np.ndarray)):
                value = json.dumps(value.tolist() if isinstance(value, np.ndarray) else value)
            elif pd.isna(value):
                value = None
            elif isinstance(value, np.generic):
                value = value.item()
            c = ws.cell(i, j, value)
            EXPECTED_CELLS.append((ws.title, i, j, value))
            if isinstance(value, str):
                c.data_type = 's'  # Preserve source formulas as text, never evaluate them in this sheet.
            elif isinstance(value, float):
                c.number_format = '0.000000E+00' if value and abs(value) < .001 else '0.000000'
            c.font = Font(name='Aptos', size=10)
    COVERAGE.append(dict(sheet=ws.title, block=title, first_data_row=header + 1,
                         rows=len(table), source=str(source), note=note))


def notebook_table(rel):
    genomic_repo = ROOT
    index = json.loads((genomic_repo / 'data/genomics/generated_tables/index.json').read_text())
    item = index[rel]
    table = pd.read_json(genomic_repo / item['file'], orient='table', precise_float=True)
    if len(table) != item['rows']:
        raise ValueError(f"Genomic table row count differs: {rel}")
    return table, item


def main():
    DEST.mkdir(exist_ok=True)
    wb = Workbook()
    wb.active.title = 'Read me'
    manifest = panel_manifest()
    figures = [f'Figure {i}' for i in range(1, 7)] + [f'Supplementary Figure {i}' for i in range(1, 14)]
    for fig in figures + [f'Supplementary Table {i}' for i in range(1, 5)]:
        wb.create_sheet(sheet_name(fig))
    wb.create_sheet('Sensitivity')
    final = pd.read_csv(HERE / 'out/final_statistics.csv')
    groups = pd.read_csv(HERE / 'out/group_summaries.csv')
    for fig, letters, ds, value, stem, drugs, stat_panel in FIT_PANELS:
        ws = wb[sheet_name(fig)]
        gf = split_strain(pd.read_pickle(HERE / 'cache' / f'{ds}.pkl'))
        cols = ['Strain', 'group', 'culturenumber', 'Antibiotic', value]
        note = ('Three technical series from one overnight culture per strain; biological n=1.'
                if ds == 'fig5_mutants' else 'One fitted estimate per culture; matched lineage IDs retained.')
        if ds == 'supp3_pcr':
            gf['included_in_comparison'] = gf.index.isin(comparison_cohort(ds, gf).index)
            cols.append('included_in_comparison')
            note += ' Paired comparison and plot use P1-P6/PCr1-PCr6. P7-P10 have no evolved partners here and remain source records only.'
        block(ws, f'Panel {"/".join(dict.fromkeys(letters))}: {value} per culture or technical series',
              gf[cols], f'cache/{ds}.pkl', note)
        if stat_panel:
            sub = final[(final.Panel == stat_panel) & (final.value == value)]
            block(ws, f'{value} contrasts', sub, 'out/final_statistics.csv',
                  'Pointwise 95% t-based ratio intervals estimate uncertainty; they are not multiplicity-adjusted. '
                  + ('MIC has no hypothesis tests, adjusted p-values, or significance calls.' if value == 'MIC'
                     else 'IC50 p-values use Holm adjustment within panel and drug.'))
            sub = groups[(groups.Panel == stat_panel) & (groups.value == value)]
            block(ws, f'{value} group summaries', sub, 'out/group_summaries.csv')
        else:
            sub = pd.read_csv(HERE / 'out/fig5_descriptive.csv').query('value == @value')
            block(ws, f'{value} descriptive summaries', sub, 'out/fig5_descriptive.csv')
    for fig, (panel, path, stem, note) in MDK.items():
        ws = wb[sheet_name(fig)]
        for suffix in ['records', 'summary']:
            file = ('supp11_recomputed.csv' if suffix == 'records' else 'supp11_summary.csv') if fig == 'Supplementary Figure 11' else f'{stem}_{suffix}.csv'
            block(ws, f'Panel {panel}: MDK {suffix}', pd.read_csv(HERE / 'out' / file), 'out/' + file,
                  note + ' Missing counts are distinct from recorded zero counts; MAD is normal-scaled.')
    for filename in ['fig5a_doubling_times_per_well.csv', 'fig5a_descriptive.csv']:
        block(wb['Figure 5'], 'Panel a: doubling time', pd.read_csv(HERE / 'out' / filename),
              'out/' + filename, 'Minutes; six technical wells of one biological culture per strain.')
    for fig in ['Figure 1', 'Figure 3', 'Supplementary Figure 3', 'Supplementary Figure 4']:
        rows = manifest[(manifest.figure == fig) & manifest.measure.isin(['survival trajectory', 'resistance trajectory'])]
        for source in rows.source.unique():
            book = pd.ExcelFile(ROOT / source)
            for sheet in book.sheet_names:
                table = pd.read_excel(book, sheet_name=sheet)
                shown = {'Figure 1': 'PA, PL, PC', 'Figure 3': 'PLAC',
                         'Supplementary Figure 3': 'PCr-evolution',
                         'Supplementary Figure 4': 'ATEC'}[fig]
                block(wb[sheet_name(fig)], f'Trajectory records, original sheet {sheet}', table, source,
                      ('Complete source workbook retained, including sheets outside this figure. '
                       if len(book.sheet_names) > 1 else 'Source workbook values retained. ') +
                      f'{fig} uses {shown}. Original sheet, strain, culture, and day identify observations.')
    from survival import records as survival_records
    survival = survival_records()
    block(wb['Figure 1'], 'Panel f: paired survival percentages', survival,
          'working/figures/Figure 1/D-F/SurvivalData.xlsx, PC',
          'Pre/post observations matched by culture and day; includes recovered culture 10 at day 20.')
    block(wb['Figure 1'], 'Panel f: daily mean and sample SD',
          survival.groupby('day').survival_percent.agg(n='count', mean='mean', sd='std').reset_index(),
          'Panel f paired survival percentages')
    block(wb['Figure 1'],'Panels d–f: paired survival percentages',pd.read_csv(HERE/'out/fig1_all_survival_records.csv'),'out/fig1_all_survival_records.csv')
    raw = pd.read_pickle(HERE / 'cache/fig2_paplpc__df_analysis.pkl')
    example = raw[(raw.Strain == 'PA5') & (raw.Antibiotic == 'Levofloxacin')]
    block(wb['Figure 2'], 'Panel a: representative PA5 levofloxacin dose response', example,
          'cache/fig2_paplpc__df_analysis.pkl; original notebook cell 29')
    fit = pd.read_pickle(HERE / 'cache/fig2_paplpc.pkl')
    fit = fit[(fit.Strain == 'PA5') & (fit.Antibiotic == 'Levofloxacin')].iloc[0]
    x_fit = json.loads(fit.x_fit) if isinstance(fit.x_fit, str) else fit.x_fit
    y_fit = json.loads(fit.y_fit) if isinstance(fit.y_fit, str) else fit.y_fit
    block(wb['Figure 2'], 'Panel a: fitted curve', pd.DataFrame({'concentration': x_fit, 'fitted_OD': y_fit}),
          'cache/fig2_paplpc.pkl')

    recovered = {}
    for fig in figures:
        ws = wb[sheet_name(fig)]
        rows = manifest[(manifest.figure == fig) & (manifest.measure == 'mutation frequency')]
        for source in rows.source.unique():
            result = recovered.setdefault(source, notebook_table(source))
            panels = ','.join(rows[rows.source == source].panel)
            table, item = result
            note = (f"Regenerated from mutation-call TSVs using notebook cells {item['cells']}; "
                    f"{len(table)} rows. Original two-decimal frequency rounding retained; "
                    "full-precision calls and input hashes are in the Lyon_2026 repository.")
            block(ws, f'Panel {panels}: regenerated mutation-frequency table', table,
                  item['file'] + '; notebook: ' + source, note)

    genomic_repo = ROOT
    crosswalk = pd.read_csv(genomic_repo/'data/genomics/plot_tables/panel_source_crosswalk.csv')
    figure_map={'Fig2':'Figure 2','Fig4':'Figure 4','S3':'Supp Fig 3','S4':'Supp Fig 4','S5':'Supp Fig 5','S6':'Supp Fig 6','S9':'Supp Fig 9'}
    for prefix,sheet in figure_map.items():
        selected=crosswalk[crosswalk.panel.str.startswith(prefix)]
        block(wb[sheet], 'Plotted mutation frequencies and source crosswalk', selected,
              'data/genomics/plot_tables/panel_source_crosswalk.csv',
              'Distinct mutation identities are separate; short labels are not averaged. Coordinates refer to the source reference, including contig IDs for PbEc.')
    block(wb['Figure 4'], 'Mutation trajectories used in Figure 4 and Supplementary Figure 8',
          pd.read_csv(genomic_repo/'data/genomics/plot_tables/PLAC_trajectories.csv'),
          'data/genomics/plot_tables/PLAC_trajectories.csv')
    block(wb['Sensitivity'],'Figure 3e possible handling failure',pd.read_csv(HERE/'out/Fig3e_handling_sensitivity.csv'),
          'out/Fig3e_handling_sensitivity.csv','Primary analysis retains the confirmed zero. Omission is a sensitivity analysis because pellet loss is not confirmed.')
    tables = json.loads(SOURCE.read_text())
    for i, rows in enumerate(tables[:2], 1):
        block(wb[f'Supp Table {i}'], f'Supplementary Table {i}', pd.DataFrame(rows[1:], columns=rows[0]),
              str(SOURCE.relative_to(ROOT)), 'Table values and column headings from the source manuscript.')
    single = genomic_repo/'data/single-cell'
    mapping = json.loads((single/'plot_tables/table_mapping.json').read_text())
    for number,sheets in mapping.items():
        for name in sheets:
            frame=pd.read_excel(genomic_repo/f'working/source-data/Supplementary Table {number}.xlsx',sheet_name=name,header=None)
            frame.columns=[f'Original column {i+1}' for i in range(frame.shape[1])]
            source = ('data/single-cell/original-tables/Single Cell Analysis Supp Tables.xlsx; '
                      f'{name}; via working/source-data/Supplementary Table {number}.xlsx')
            note = 'Original worksheet cells retained.'
            if name == 'up in WT vs 4 and 7 GSEA':
                source += '; data/single-cell/enrichment/exports/Enriched-from-Erdal-up-in-WT-vs-common-4-and-7.txt'
                note += ' All 133 export rows are included; the original 19 are the P < 0.001 subset. See data/single-cell/enrichment/README.md.'
            block(wb[f'Supp Table {number}'],name,frame,source,note)
    for name,title in [('sample_metadata.csv','Technical sample metadata'),('cluster_composition.csv','Panel c: cluster composition'),('figure6_volcano_points.csv','Panels d/e: volcano points')]:
        block(wb['Figure 6'],title,pd.read_csv(single/'plot_tables'/name),'data/single-cell/plot_tables/'+name)
    cells=pd.read_csv(single/'generated/cells_and_embeddings.csv')
    expression=pd.read_csv(single/'generated/figure6_expression.csv')
    cells=cells[['barcode','sample','technical_sample','cluster_published','umap_1','umap_2']].merge(expression,on='barcode',validate='one_to_one')
    block(wb['Figure 6'],'Panels a/b/f/g: cell embeddings and normalized expression',cells,
          'data/single-cell/generated/cells_and_embeddings.csv; figure6_expression.csv',
          '48,883 cells from six technical samples, two per culture. Each technical sample was probed and processed through microfluidics separately.')
    targeted=pd.concat([pd.read_csv(single/f'generated/marker_claim_check_{culture}.csv') for culture in [4,7]])
    block(wb['Supp Table 4'],'Targeted marker checks',targeted,'data/single-cell/generated/marker_claim_check_4.csv; marker_claim_check_7.csv','Targeted hipA and rplJ comparisons use min.pct=0 and logfc.threshold=0, with Bonferroni correction over all assay features.')
    points=pd.read_csv(single/'plot_tables/supplementary_volcano_points.csv')
    for number,clusters in [(12,[1,2]),(13,[4,9])]:
        block(wb[f'Supp Fig {number}'],'Volcano points and original chart annotations',points[points.cluster.isin(clusters)],
              'data/single-cell/plot_tables/supplementary_volcano_points.csv',
              'Every point matched uniquely to the recomputed full DGE table. Colors retained from the supplied Excel chart point annotations; zero P-values capped at 10^-304 for display.')
    for filename in ['sensitivity_all.csv', 'what_changes_vs_published.csv', 'what_changes_vs_17aug_report.csv', 'experimental_units.csv']:
        block(wb['Sensitivity'], filename, pd.read_csv(HERE / 'out' / filename), 'out/' + filename,
              'Sensitivity p-values identify their family. Fit-perturbation bounds are empirical sensitivity ranges.')
    block(wb['Read me'], 'Source Data', manifest, 'out/panel_manifest.csv',
          'Censored MDK records and single-cell technical sample identities are identified in their sheets. Genomic tables are regenerated from the supplied mutation calls. '
          'One sheet per figure; full-precision numerical cells and source paths retained.')
    for ws in wb:
        ws.freeze_panes = 'C6'
        ws.sheet_view.zoomScale = 85
        for col in range(1, ws.max_column + 1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col)].width = 23 if col < 4 else 19
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    # Full annotations are longer than notebook display snippets. Wrap the new
    # genomic blocks without changing the layout of preceding source records.
    import math
    from openpyxl.utils import get_column_letter
    for entry in COVERAGE:
        if 'regenerated mutation-frequency table' not in entry['block']:
            continue
        ws = wb[entry['sheet']]
        first = entry['first_data_row']
        ws.row_dimensions[first - 1].height = 30
        for row in ws.iter_rows(min_row=first, max_row=first + entry['rows'] - 1):
            lines = 1
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if isinstance(cell.value, str):
                    width = ws.column_dimensions[get_column_letter(cell.column)].width
                    lines = max(lines, math.ceil(len(cell.value) / (width * 0.9)))
            ws.row_dimensions[row[0].row].height = min(409, 14 * lines + 4)
    file = DEST / 'Source Data.xlsx'
    wb.save(file)
    check = load_workbook(file, data_only=False, read_only=False)
    if len(check.sheetnames) != 25:
        raise ValueError(f'Unexpected sheet count: {len(check.sheetnames)}')
    numeric = text = blanks = 0
    for sheet, row, col, expected in EXPECTED_CELLS:
        cell = check[sheet].cell(row, col)
        actual = cell.value
        if expected is None or expected == '':
            assert actual is None or actual == '', (sheet, row, col, expected, actual)
            blanks += 1
        elif isinstance(expected, (float, int)):
            assert np.isclose(actual, expected, rtol=1e-14, atol=1e-300), (sheet, row, col, expected, actual)
            numeric += 1
        else:
            assert actual == expected, (sheet, row, col, expected, actual)
            assert cell.data_type != 'f', (sheet, row, col)
            text += 1
    (DEST / 'source_data_validation.json').write_text(json.dumps(dict(
        sheets=len(check.sheetnames), blocks=len(COVERAGE), numeric_cells=numeric,
        text_cells=text, blank_cells=blanks, all_cells_match_source_tables=True), indent=2)+'\n')
    pd.DataFrame(COVERAGE).to_csv(DEST / 'source_data_blocks.csv', index=False)
    print(f'{len(check.sheetnames)} sheets; {len(COVERAGE)} source-data blocks; {file}')


if __name__ == '__main__':
    main()
