"""Validate the supplied single-cell analysis and render its numerical figures."""
from pathlib import Path
from zipfile import ZipFile
import hashlib,json,os
import xml.etree.ElementTree as E
os.environ.setdefault('MPLCONFIGDIR','/Users/wak/micromamba/mpl_cache')
os.environ['MPLBACKEND']='Agg'
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'data/single-cell'; GEN=BASE/'generated'; TABLES=BASE/'original-tables'
OUT=ROOT/'working/figures-assembled'; EXPORT=BASE/'plot_tables'
BOOK=TABLES/'Single Cell Analysis Supp Tables.xlsx'
COLORS=['#97269e','#a4cb3b','#356abb','#dba03b','#50b7ce','#d77832','#4b9986','#d83334','#42a342']
CHECKS={}

def compare(a,b,keys):
 a=a.set_index(keys).sort_index();b=b.set_index(keys).sort_index()
 assert a.index.equals(b.index),'Differential-expression row identities differ'
 report={'rows':len(a)}
 for col in ['avg_log2FC','pct.1','pct.2']:
  assert np.allclose(a[col],b[col],rtol=1e-12,atol=1e-12),col
  report[col+'_max_absolute_difference']=float(np.abs(a[col]-b[col]).max())
 for col in ['p_val','p_val_adj']:
  assert np.array_equal(a[col]==0,b[col]==0),(col,'zero pattern')
  m=a[col]>0;delta=float(np.abs(np.log10(a.loc[m,col])-np.log10(b.loc[m,col])).max())
  assert delta<1e-7,(col,delta)
  report[col+'_max_log10_difference']=delta
 return report

def verify():
 manifest=json.loads((BASE/'input_manifest.json').read_text())
 # The archive manifest is preserved; verify every selected input independently.
 entries=manifest['files']
 for item in entries:
  p=BASE/item['path'];assert hashlib.sha256(p.read_bytes()).hexdigest()==item['sha256'],str(p)
 cells=pd.read_csv(GEN/'cells_and_embeddings.csv')
 expected=[9335,7034,6008,5989,5092,4967,4961,3052,2445]
 assert cells.cluster_published.value_counts().sort_index().tolist()==expected
 assert cells.technical_sample.value_counts().sort_index().tolist()==[7827,10774,4764,2764,9619,13135]
 assert cells.barcode.is_unique and len(cells)==48883
 CHECKS['cells']=len(cells);CHECKS['cluster_sizes']=expected;CHECKS['input_files_verified']=len(entries)
 for file,sheet,keys in [('cluster_published_DGE_recomputed.csv','cluster_DGE',['cluster','gene']),('sample4_vs_parent_recomputed.csv','DGE 4 vs parent',['gene']),('sample7_vs_parent_recomputed.csv','DGE 7 vs parent',['gene'])]:
  CHECKS[sheet]=compare(pd.read_csv(GEN/file),pd.read_excel(BOOK,sheet_name=sheet),keys)
 CHECKS['sample_DGE']=compare(pd.read_csv(GEN/'sample_DGE_recomputed.csv'),pd.read_csv(TABLES/'sample_DGE_full.csv'),['cluster','gene'])
 return cells

def export_tables():
 mapping={3:['Sample statistics','cluster_DGE','up in 1 and gsea w benj hoch','up seurat clstr2 and gsea ','up seurat clstr9 and gsea','up seurat clstr4 and gsea'],4:['DGE 4 vs parent','DGE 7 vs parent','up in 4 and 7 vs WT GSEA','up in WT vs 4 and 7 GSEA']}
 dest=ROOT/'working/source-data';dest.mkdir(exist_ok=True)
 for number,sheets in mapping.items():
  wb=load_workbook(BOOK)
  for sheet in list(wb.sheetnames):
   if sheet not in sheets:del wb[sheet]
  assert wb.sheetnames==sheets
  if number==4:
   sheet=wb.create_sheet('Targeted marker checks')
   targeted=pd.concat([pd.read_csv(GEN/f'marker_claim_check_{culture}.csv') for culture in [4,7]])
   sheet.append(['Targeted checks of hipA and rplJ with min.pct=0 and logfc.threshold=0; Bonferroni correction uses all assay features.'])
   sheet.append(list(targeted.columns))
   for row in targeted.itertuples(index=False,name=None):sheet.append(list(row))
  path=dest/f'Supplementary Table {number}.xlsx';wb.save(path)
  original=load_workbook(BOOK,data_only=False);check=load_workbook(path,data_only=False)
  for sheet in sheets:
   for row_a,row_b in zip(original[sheet].values,check[sheet].values):
    for a,b in zip(row_a,row_b):
     if isinstance(a,(float,int)):assert np.isclose(a,b,rtol=1e-14,atol=1e-300),(sheet,a,b)
     else:assert a==b,(sheet,a,b)
  CHECKS[f'supplementary_table_{number}']={'sheets':sheets,'original_cells_preserved':True}
 (EXPORT/'table_mapping.json').write_text(json.dumps(mapping,indent=2)+'\n')
 samples=pd.read_csv(GEN/'filtering_counts.csv').rename(columns={'sample':'technical_sample'})
 samples['culture']=samples.technical_sample.str[0].replace({'3':'parent'})
 samples['replicate']=samples.technical_sample.str[-1].map({'a':1,'b':2})
 samples['replication']='same culture sampling; separately probed and microfluidically processed'
 samples.to_csv(EXPORT/'sample_metadata.csv',index=False)
 counts=pd.read_csv(GEN/'cluster_sample_counts.csv');counts['percent_within_cluster']=counts.Freq/counts.groupby('cluster').Freq.transform('sum')*100
 counts.to_csv(EXPORT/'cluster_composition.csv',index=False)


def supplement_points():
 book=TABLES/'volcanos cutoff supp figures.xlsx';all_dge=pd.read_csv(GEN/'cluster_published_DGE_recomputed.csv');rows=[]
 ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
 with ZipFile(book) as z:
  for chart,cluster in enumerate([1,2,4,9],1):
   frame=pd.read_excel(book,sheet_name=f'cluster {cluster} volcano');sub=all_dge[all_dge.cluster==cluster]
   tree=E.fromstring(z.read(f'xl/charts/chart{chart}.xml'));styles={}
   for point in tree.findall('.//c:dPt',ns):
    idx=int(point.find('c:idx',ns).get('val'));fill=point.find('c:marker/c:spPr/a:solidFill',ns)
    if fill is not None:
     color=list(fill)[0].get('val');styles[idx]={'FF0000':'red','7030A0':'purple','tx1':'black'}.get(color,color)
   for idx,row in frame.dropna(subset=['gene']).iterrows():
    matched=sub[np.isclose(sub.avg_log2FC,float(row.avg_log2FC),rtol=1e-10,atol=1e-12)]
    assert len(matched)==1,(cluster,row.gene,len(matched))
    r=matched.iloc[0]
    assert abs(r.avg_log2FC)>.25 and r.p_val_adj<.05
    rows.append(dict(cluster=cluster,source_excel_row=idx+2,display_gene=row.gene,probe=r.gene,avg_log2FC=r.avg_log2FC,p_val_adj=r.p_val_adj,neg_log10_p_adj=-np.log10(max(r.p_val_adj,1e-304)),color=styles.get(idx,'other')))
 points=pd.DataFrame(rows);points.to_csv(EXPORT/'supplementary_volcano_points.csv',index=False)
 CHECKS['supplementary_points']=points.groupby('cluster').size().to_dict()
 # Every plotted point maps uniquely to the full recomputed cluster table.
 return points


def style(ax):
 ax.tick_params(labelsize=8.5,width=.65,length=2.5,pad=2)
 for spine in ax.spines.values():spine.set_linewidth(.65)
 ax.spines[['top','right']].set_visible(False)

def save(fig,name):
 OUT.mkdir(exist_ok=True);fig.savefig(OUT/f'{name}.pdf',metadata={'CreationDate':None,'ModDate':None})
 fig.savefig(OUT/f'{name}.png',dpi=300);fig.savefig(OUT/f'{name}_preview.png',dpi=100);plt.close(fig)

def singlecell_figure(cells):
 fig=plt.figure(figsize=(6.45,7.4));fig.text(.02,.975,'Figure 6',weight='bold',fontsize=12)
 x,y=cells.umap_1.to_numpy(),cells.umap_2.to_numpy();cluster=cells.cluster_published.to_numpy()
 limits=[(x.min()-.5,x.max()+.5),(y.min()-.5,y.max()+.5)]
 def umap(rect,mask=None,expression=None):
  ax=fig.add_axes(rect);sel=np.ones(len(cells),dtype=bool) if mask is None else mask
  if expression is None:
   for k,color in enumerate(COLORS,1):
    take=sel&(cluster==k);ax.scatter(x[take],y[take],s=.4,c=color,linewidths=0,rasterized=True)
  else:
   order=np.argsort(expression);ax.scatter(x[order],y[order],s=.4,c=expression[order],cmap='Reds',vmin=0,vmax=7,linewidths=0,rasterized=True)
  ax.set_xlim(*limits[0]);ax.set_ylim(*limits[1]);ax.set_xticks([]);ax.set_yticks([]);ax.set_aspect('equal');style(ax)
  return ax
 a=umap([.07,.695,.30,.235]);fig.text(.02,.935,'a',weight='bold',fontsize=12)
 for k in range(1,10):
  m=cluster==k;a.text(np.median(x[m]),np.median(y[m]),str(k),fontsize=8.5,weight='bold',ha='center',va='center',bbox=dict(facecolor='white',alpha=.85,edgecolor='none',pad=.5))
 a.set_xlabel('UMAP 1',fontsize=9.5);a.set_ylabel('UMAP 2',fontsize=9.5)
 fig.text(.42,.935,'b',weight='bold',fontsize=12)
 for i,(sample,title) in enumerate([('parent','Parent'),('4','Culture 4'),('7','Culture 7')]):
  ax=umap([.44+i*.185,.695,.17,.235],cells['sample'].astype(str).eq(sample).to_numpy());ax.set_title(title,fontsize=9.5,pad=5)
 c=fig.add_axes([.10,.41,.30,.19]);fig.text(.02,.62,'c',weight='bold',fontsize=12)
 counts=pd.read_csv(EXPORT/'cluster_composition.csv');bottom=np.zeros(9)
 for sample,color,label in [('parent','.55','Parent'),('4','#238b45','Culture 4'),('7','#c03a73','Culture 7')]:
  v=counts[counts['sample'].astype(str)==sample].sort_values('cluster').percent_within_cluster.to_numpy();c.bar(range(1,10),v,bottom=bottom,color=color,width=.8,label=label);bottom+=v
 c.set_xticks(range(1,10));c.set_yticks([0,50,100]);c.set_xlabel('Cluster',fontsize=9.5);c.set_ylabel('Cells in cluster (%)',fontsize=9.5);style(c)
 c.legend(loc='lower left',bbox_to_anchor=(-.1,1.06),ncol=3,fontsize=7.5,frameon=False,handlelength=.7,columnspacing=.7)
 expr=pd.read_csv(GEN/'figure6_expression.csv').set_index('barcode').loc[cells.barcode]
 for i,(gene,probe) in enumerate([('rplJ','rplJ-2'),('hipA','hipA-3')]):
  ax=umap([.49+i*.25,.405,.20,.20],expression=expr[probe].to_numpy());ax.set_title(gene,fontsize=10,fontstyle='italic',pad=4);fig.text(.45+i*.25,.62,['f','g'][i],weight='bold',fontsize=12)
 from matplotlib.cm import ScalarMappable
 from matplotlib.colors import Normalize
 cb=fig.colorbar(ScalarMappable(norm=Normalize(0,7),cmap='Reds'),cax=fig.add_axes([.55,.375,.35,.012]),orientation='horizontal',ticks=[0,3.5,7]);cb.ax.tick_params(labelsize=8,length=2,width=.6);cb.outline.set_linewidth(.5)
 cb.set_label('Log-normalized expression',fontsize=8.5,labelpad=1)
 plotted=[]
 for i,culture in enumerate([4,7]):
  d=pd.read_csv(GEN/f'sample{culture}_vs_parent_recomputed.csv');source=pd.read_excel(BOOK,sheet_name=f'DGE {culture} vs parent').set_index('gene')
  d['shared_status']=source.loc[d.gene,'shared_status'].to_numpy();d=d[(d.avg_log2FC.abs()>.5)&(d.p_val_adj<.05)].copy();d['neg_log10_p_adj']=-np.log10(d.p_val_adj.clip(lower=1e-304));d['culture']=culture
  shared=d.shared_status.str.lower().str.contains('shared')
  d['color']=np.where(shared,np.where(d.avg_log2FC>0,'#2369a0','#cf2b35'),'.7')
  ax=fig.add_axes([.10+i*.50,.09,.35,.205]);fig.text(.02+i*.50,.315,['d','e'][i],weight='bold',fontsize=12)
  for color,g in d.groupby('color'):ax.scatter(g.avg_log2FC,g.neg_log10_p_adj,c=color,s=6,linewidths=0,rasterized=False)
  ax.axvline(0,color='.7',lw=.5);ax.set_ylim(0,320);ax.set_yticks([0,100,200,300]);ax.set_xlabel(r'Log$_2$ fold change',fontsize=9.5);ax.set_ylabel(r'$-\log_{10}$ adjusted P',fontsize=9.5);ax.set_title(f'Culture {culture} versus parent',fontsize=9.5,pad=5);style(ax);plotted.append(d)
 fig.legend(handles=[Line2D([],[],marker='o',ls='',ms=3,color='#2369a0',label='Higher in both evolved cultures'),Line2D([],[],marker='o',ls='',ms=3,color='#cf2b35',label='Higher in parent in both comparisons')],loc='lower center',bbox_to_anchor=(.51,.008),ncol=1,fontsize=8,frameon=False,labelspacing=.3)
 pd.concat(plotted).to_csv(EXPORT/'figure6_volcano_points.csv',index=False);save(fig,'Figure_6')


def supplementary_figures(points):
 for number,clusters in [(12,[1,2]),(13,[4,9])]:
  fig,axes=plt.subplots(2,1,figsize=(6.45,6.6));fig.subplots_adjust(left=.13,right=.97,bottom=.085,top=.83,hspace=.43)
  fig.text(.02,.97,f'Supplementary Figure {number}',fontsize=12,weight='bold')
  for ax,cluster in zip(axes,clusters):
   sub=points[points.cluster==cluster]
   for color,g in sub.groupby('color',sort=False):
    ax.scatter(g.avg_log2FC,g.neg_log10_p_adj,s=11 if color!='other' else 7,facecolors={'other':'none','red':'#df292f','black':'.1','purple':'#7030A0'}[color],edgecolors='#65b6d0' if color=='other' else 'none',linewidths=.45)
   ax.set_title(f'Cluster {cluster}',loc='left',fontsize=10);ax.axvline(0,lw=.5,color='.7');ax.set_xlabel(r'Log$_2$ fold change',fontsize=10);ax.set_ylabel(r'$-\log_{10}$ adjusted P',fontsize=10);ax.set_ylim(bottom=0);style(ax)
  labels=[('red','Translation'),('black','Fimbriae')] if number==12 else [('red','Purine nucleotide binding'),('purple','Glycolysis, pyruvate dehydrogenase, TCA and glyoxylate bypass')]
  fig.legend(handles=[Line2D([],[],marker='o',ls='',ms=4,color={'red':'#df292f','black':'.1','purple':'#7030A0'}[c],label=t) for c,t in labels],loc='upper left',bbox_to_anchor=(.115,.952),ncol=1,fontsize=8.5,frameon=False,labelspacing=.35)
  save(fig,f'Supplementary_Figure_{number}')


def main():
 EXPORT.mkdir(exist_ok=True);plt.rcParams.update({'font.family':'Times New Roman','font.size':9.5,'pdf.fonttype':42,'svg.fonttype':'none','mathtext.fontset':'stix'})
 cells=verify();export_tables();points=supplement_points();singlecell_figure(cells);supplementary_figures(points)
 CHECKS['enrichment']='Supplied result tables preserved; generating method, database version and tested background are not in the supplied Rmd files.'
 (EXPORT/'validation.json').write_text(json.dumps(CHECKS,indent=2)+'\n');print(json.dumps(CHECKS,indent=2))
if __name__=='__main__':main()
