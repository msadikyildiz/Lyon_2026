"""Verify original EcoCyc exports and recover the complete parent-up result table."""
import csv
from copy import copy
import hashlib
import json
import math
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'data/single-cell/enrichment'
BOOK = BASE.parent / 'original-tables/Single Cell Analysis Supp Tables.xlsx'


def exports():
    items = json.loads((BASE / 'manifest.json').read_text())['files']
    for item in items:
        p = BASE / item['path']
        data = p.read_bytes()
        assert len(data) == item['bytes'] and hashlib.sha256(data).hexdigest() == item['sha256'], str(p)
        with p.open() as stream:
            rows = list(csv.reader(stream, delimiter='\t'))
        assert rows[0][1:] == ['p-values', 'Matches'], p.name
        assert len(rows) - 1 == item['rows'] and all(len(row) == 3 for row in rows), p.name
        yield item, rows[1:]


def verify():
    wb = load_workbook(BOOK, data_only=True)
    report = []
    for item, rows in exports():
        matched = 0
        if item['worksheet']:
            sheet = wb[item['worksheet']]
            col = item['first_column']
            for index, row in enumerate(rows, item['first_data_row']):
                values = [sheet.cell(index, c).value for c in range(col, col + 3)]
                if values == [None, None, None]:
                    assert item['worksheet'] == 'up in WT vs 4 and 7 GSEA' and index >= 21
                    continue
                assert row[0] == values[0], (item['worksheet'], index)
                assert math.isclose(float(row[1]), values[1], rel_tol=1e-12, abs_tol=1e-300)
                assert (row[2] or None) == values[2], (item['worksheet'], index, 'gene list')
                matched += 1
            # An export must not silently omit any existing result row.
            existing = sum(sheet.cell(i, col).value is not None for i in range(item['first_data_row'], sheet.max_row + 1))
            assert matched == existing
        report.append({'export': item['path'], 'rows': len(rows), 'worksheet': item['worksheet'],
                       'existing_rows_matched': matched,
                       'rows_added': len(rows) - matched if item['worksheet'] else 0})
    result = {'exports_verified': len(report), 'comparisons': report,
              'original_workbook_unchanged': True, 'enrichment_recalculated': False}
    (BASE / 'validation.json').write_text(json.dumps(result, indent=2) + '\n')
    return result


def extend_parent_results(wb):
    """Append missing export rows, preserving every existing workbook cell."""
    sheet = wb['up in WT vs 4 and 7 GSEA']
    assert sheet.max_row == 20
    for item, rows in exports():
        if item['worksheet'] != sheet.title:
            continue
        for row_index, row in enumerate(rows[19:], 21):
            for col in range(1, 13):
                sheet.cell(row_index, col)._style = copy(sheet.cell(20, col)._style)
            for col, value in zip([10, 11, 12], [row[0], float(row[1]), row[2] or None]):
                sheet.cell(row_index, col).value = value
        sheet['J1'].comment = Comment(
            'All 133 rows from the matching original EcoCyc export are retained. '
            'The first 19 match the original workbook; rows 21–134 restore the remaining results. '
            'Exported P-values extend to <0.1, not all <0.05. See data/single-cell/enrichment/README.md.',
            'Muhammed Sadik Yildiz')


if __name__ == '__main__':
    print(json.dumps(verify(), indent=2))
