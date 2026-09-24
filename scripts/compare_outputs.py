"""Check numerical outputs and report rendering differences across environments."""
import argparse
import hashlib
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
RTOL, ATOL = 1e-8, 1e-10
OUT = Path('working/analysis/stats-rework/out')

def compare(reference, rebuilt, strict_images=False):
    if reference.resolve() == rebuilt.resolve():
        raise ValueError('Reference and rebuilt directories must differ')
    names = {p.name for p in (reference/OUT).glob('*.csv')}
    if not names or names != {p.name for p in (rebuilt/OUT).glob('*.csv')}:
        raise ValueError('CSV output file set differs or is empty')
    for name in sorted(names):
        a, b = pd.read_csv(reference/OUT/name), pd.read_csv(rebuilt/OUT/name)
        if name == 'fig5a_descriptive.csv':
            for x, y in zip(a['values'], b['values']):
                np.testing.assert_allclose(list(map(float,x.split(';'))), list(map(float,y.split(';'))), rtol=RTOL, atol=ATOL)
            a, b = a.drop(columns='values'), b.drop(columns='values')
        pd.testing.assert_frame_equal(a,b,check_exact=False,rtol=RTOL,atol=ATOL,obj=name)
    pngs, image_differences = {}, []
    for relative in [OUT/'figures-final', OUT/'print-panels', OUT/'trajectories', OUT/'genomic-panels', Path('working/figures-assembled')]:
        files = {p.name for p in (reference/relative).glob('*.png')}
        if not files or files != {p.name for p in (rebuilt/relative).glob('*.png')}:
            raise ValueError(f'PNG file set differs or is empty: {relative}')
        for name in sorted(files):
            a, b = reference/relative/name, rebuilt/relative/name
            with Image.open(a) as ai, Image.open(b) as bi:
                ai.verify(); bi.verify()
                if ai.size != bi.size:
                    raise ValueError(f'PNG dimensions differ: {relative/name}')
            if hashlib.sha256(a.read_bytes()).digest() != hashlib.sha256(b.read_bytes()).digest():
                image_differences.append((relative/name).as_posix())
        pngs[str(relative)] = len(files)
    if strict_images and image_differences:
        raise ValueError(f'PNG differs: {image_differences[0]}')
    generated_files={}
    for folder in ['data/genomics/generated_tables','data/genomics/plot_tables','data/single-cell/plot_tables']:
        relative=Path(folder)
        files={p.name for p in (reference/relative).iterdir() if p.is_file()}
        assert files and files=={p.name for p in (rebuilt/relative).iterdir() if p.is_file()},folder
        for name in files:
            a, b = reference/relative/name, rebuilt/relative/name
            if folder == 'data/genomics/plot_tables' and name == 'manifest.json':
                # Script versions and byte hashes are provenance, not numerical results.
                expected, actual = json.loads(a.read_text()), json.loads(b.read_text())
                assert set(expected['files']) == set(actual['files'])
                for path, digest in actual.pop('files').items():
                    assert hashlib.sha256((rebuilt/path).read_bytes()).hexdigest() == digest, path
                assert hashlib.sha256((rebuilt/'scripts/reproduce_genomic_plots.py').read_bytes()).hexdigest() == actual.pop('plot_script_sha256')
                expected.pop('files'); expected.pop('plot_script_sha256')
                assert expected == actual, str(relative/name)
            else:
                assert a.read_text(encoding='utf-8') == b.read_text(encoding='utf-8'), str(relative/name)
        generated_files[folder]=len(files)
    a=load_workbook(reference/'working/source-data/Source Data.xlsx',read_only=True,data_only=False)
    b=load_workbook(rebuilt/'working/source-data/Source Data.xlsx',read_only=True,data_only=False)
    if a.sheetnames != b.sheetnames:
        raise ValueError('Workbook sheets differ')
    numeric = cells = 0
    for name in a.sheetnames:
        if (a[name].max_row,a[name].max_column)!=(b[name].max_row,b[name].max_column):
            raise ValueError(f'Sheet dimensions differ: {name}')
        for ar,br in zip(a[name],b[name]):
            for x,y in zip(ar,br):
                if x.data_type != y.data_type:
                    raise ValueError(f'Cell type differs: {name}!{x.coordinate}')
                if isinstance(x.value,(int,float)):
                    if not isinstance(y.value,(int,float)) or not np.isclose(x.value,y.value,rtol=RTOL,atol=ATOL,equal_nan=True):
                        raise ValueError(f'Numeric cell differs: {name}!{x.coordinate}')
                    numeric += 1
                elif x.value != y.value:
                    # Doubling-time values are also stored as semicolon-separated numeric text.
                    try:
                        xv=list(map(float,x.value.split(';')));yv=list(map(float,y.value.split(';')))
                    except (AttributeError,ValueError):
                        raise ValueError(f'Text cell differs: {name}!{x.coordinate}') from None
                    np.testing.assert_allclose(xv,yv,rtol=RTOL,atol=ATOL)
                cells += 1
    a.close();b.close()
    return {'status':'passed','csv_tables':len(names),'pngs_checked':pngs,'pngs_different':image_differences,
            'strict_images':strict_images,'workbook_cells_checked':cells,'generated_files_checked':generated_files,
            'workbook_numeric_cells':numeric,'rtol':RTOL,'atol':ATOL,
            'comparison':'CSV values, generated table text (line endings ignored), provenance hashes, workbook values/types and PNG dimensions; PNG byte differences reported separately'}

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('reference',type=Path);parser.add_argument('rebuilt',type=Path)
    parser.add_argument('--report',type=Path)
    parser.add_argument('--strict-images',action='store_true',help='Fail on any PNG byte difference (requires matching fonts and renderer).')
    args=parser.parse_args()
    # Clear any previous success report before a comparison that may fail.
    if args.report:
        args.report.parent.mkdir(parents=True,exist_ok=True)
        args.report.write_text(json.dumps({'status':'running'})+'\n')
    try:
        report=compare(args.reference,args.rebuilt,args.strict_images)
    except Exception as error:
        if args.report:args.report.write_text(json.dumps({'status':'failed','error':str(error)},indent=2)+'\n')
        raise
    if args.report:args.report.write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))

if __name__ == '__main__':
    main()
